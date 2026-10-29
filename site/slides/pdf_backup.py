# /// script
# requires-python = ">=3.12"
# dependencies = ["playwright>=1.63", "tyro", "pillow", "python-pptx"]
# ///
"""Backups of a talk, for a room where the browser deck cannot run: every click of every slide as a page, as the deck
shows it, in a PDF and in a PowerPoint file that Google Slides opens (each page carrying the notes said over it, for
the presenter view), and a presenter's copy with each slide's last state, its timing and its speaker notes.

    node site/serve.mjs 8765 &                            # from the repository root
    uv run site/slides/pdf_backup.py --url http://127.0.0.1:8765/slides/2026-10-29-Montreal-AI-and-Neuroscience-MAIN/ \
        --out site/.screens/pdf/2026-10-29-Montreal-AI-and-Neuroscience-MAIN

Writes into --out: frames/ (shoot.py's capture of each click at --scale, with ?present&capture hiding the deck's
controls, and its errors.txt and layout.txt), slides.pdf (a 16:9 page per click, or per still where an animation names
its stills with data-stills, the appendix included), slides.pptx (the same pages, each with the notes said over it),
notes.pdf (per slide: its last state, the timing lines and the notes of deck.js's speaker view) and check.txt (what
was verified). Exits with status 1 when the capture logged a browser error, a frame is blank, or a slide has no page. Do not edit
site/ while it runs: the dev server reloads open pages on every change.
"""

import asyncio
import html
import io
import re
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path

import tyro
from PIL import Image, ImageStat
from pptx import Presentation
from pptx.util import Emu
from playwright.async_api import async_playwright

PAGE_PT = (1280, 720)  # a slide page in PDF points, the deck's own size


@dataclass(frozen=True)
class Backup:
    url: str
    out: Path
    scale: float = 2.0
    # A capture (?capture) has no transitions (deck.js), so a click lands at once: the wait is for loading when a slide
    # comes up, and a moment for each later click's images.
    wait_ms: int = 1500
    step_ms: int = 300
    jpeg_quality: int = 92


def capture(args: Backup, frames: Path) -> int:
    present = args.url + ("&" if "?" in args.url else "?") + "present&capture"
    shoot = Path(__file__).with_name("shoot.py")
    command = ["uv", "run", str(shoot), "--url", present, "--out", str(frames), "--steps",
               "--scale", str(args.scale), "--wait-ms", str(args.wait_ms), "--step-ms", str(args.step_ms)]
    with open(frames.parent / "capture.log", "w") as log:
        return subprocess.run(command, stdout=log, stderr=subprocess.STDOUT).returncode


async def slide_notes(url: str) -> list[dict]:
    """Each slide's id, title and speaker notes (HTML, with the timing lines deck.js adds), in deck order."""
    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch(channel="chrome")
        page = await browser.new_page(viewport={"width": 1280, "height": 720})
        await page.goto(url + ("&" if "?" in url else "?") + "present")
        await page.wait_for_function("window.deck?.isReady()")
        slides = await page.evaluate("""() => deck.getHorizontalSlides().map((s) => ({
          id: s.id, title: s.querySelector('h2')?.textContent.trim() ?? '',
          notes: s.querySelector('aside.notes')?.innerHTML ?? '' }))""")
        await browser.close()
        return slides


FRAME = re.compile(r"^(?P<n>\d\d)-(?P<slide>.+)-(?P<step>\d\d)(?:-s(?P<still>\d+))?$")


@dataclass(frozen=True)
class Frame:
    """An image of shoot.py --steps: a click state of a slide, or one of the stills it names for an animation."""
    path: Path
    step: int
    still: int  # 0 for a click's single image, else 1, 2, ... in capture order


def frames_by_slide(frames: Path) -> dict[str, list[Frame]]:
    """shoot.py's NN-<id>-SS.png and NN-<id>-SS-sK.png files, grouped by slide id, in deck, click and still order."""
    parsed = []
    for path in frames.glob("[0-9][0-9]-*.png"):
        match = FRAME.match(path.stem)
        if not match:
            raise ValueError(f"pdf_backup: {path.name} is not a shoot.py --steps frame name")
        parsed.append((int(match["n"]), int(match["step"]), int(match["still"] or 0), match["slide"], path))
    grouped: dict[str, list[Frame]] = {}
    for _, step, still, slide, path in sorted(parsed):
        grouped.setdefault(slide, []).append(Frame(path, step, still))
    return grouped


def slides_pdf(pages: list[Path], out: Path, quality: int) -> None:
    images = [Image.open(p).convert("RGB") for p in pages]
    dpi = images[0].width / (PAGE_PT[0] / 72)  # the frames' pixels over the page's inches
    images[0].save(out, save_all=True, append_images=images[1:], resolution=dpi, quality=quality)


def notes_by_click(notes_html: str, clicks: int) -> list[str]:
    """The notes as said over each click state: the timing lines and the lines before the first [CLICK] for the first
    state, then the lines between [CLICK] marks; all of them on the first state when the marks do not match the clicks."""
    text = lambda fragment: html.unescape(re.sub(r"<[^>]+>", "", fragment)).strip()
    timing = [text(p) for p in re.findall(r"<p>(.*?)</p>", notes_html, re.S)]
    lines = [text(li) for li in re.findall(r"<li>(.*?)</li>", notes_html, re.S)]
    parts = [[]]
    for line in lines:
        if line == "[CLICK]":
            parts.append([])
        else:
            parts[-1].append(line)
    if len(parts) != clicks:
        parts = [[line for line in lines if line != "[CLICK]"]] + [[] for _ in range(clicks - 1)]
    parts[0] = timing + parts[0]
    return ["\n".join(part) for part in parts]


def slides_pptx(slides: list[dict], grouped: dict[str, list[Frame]], out: Path, quality: int) -> None:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Emu(12192000), Emu(6858000)  # 13.333 x 7.5 in, 16:9
    blank = deck.slide_layouts[6]
    for slide in slides:
        frames = grouped[slide["id"]]
        steps = sorted({f.step for f in frames})
        by_step = dict(zip(steps, notes_by_click(slide["notes"], len(steps))))
        for i, frame in enumerate(frames):
            # A click's notes go with its first image; its further stills carry none.
            first = i == 0 or frames[i - 1].step != frame.step
            page = deck.slides.add_slide(blank)
            jpeg = io.BytesIO()
            Image.open(frame.path).convert("RGB").save(jpeg, "JPEG", quality=quality)
            jpeg.seek(0)
            page.shapes.add_picture(jpeg, 0, 0, deck.slide_width, deck.slide_height)
            page.notes_slide.notes_text_frame.text = by_step[frame.step] if first else ""
    deck.save(out)


async def notes_pdf(slides: list[dict], grouped: dict[str, list[Frame]], out: Path) -> None:
    rows = []
    for number, slide in enumerate(slides, 1):
        last = grouped[slide["id"]][-1].path
        rows.append(f"""<section><header><b>{number}. {html.escape(slide['title'])}</b></header><img src="{last.resolve().as_uri()}">
          <div class="notes">{slide['notes']}</div></section>""")
    page_html = """<!doctype html><meta charset="utf-8"><style>
      @page { size: A4; margin: 12mm; }
      body { font: 10.5pt/1.35 -apple-system, "Helvetica Neue", sans-serif; color: #0f172a; }
      section { break-after: page; }
      header { display: flex; justify-content: space-between; margin-bottom: 6px; font-size: 12pt; }
      header span { color: #64748b; font-family: ui-monospace, monospace; font-size: 9pt; }
      img { width: 100%; border: 1px solid #cbd5e1; }
      .notes p { margin: 6px 0 2px; font-family: ui-monospace, monospace; font-size: 9.5pt; color: #075985; }
      .notes ul { margin: 6px 0 0; padding-left: 18px; } .notes li { margin: 1px 0; }
    </style>""" + "".join(rows)
    source = out.with_suffix(".html")
    source.write_text(page_html)
    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch(channel="chrome")
        page = await browser.new_page()
        await page.goto(source.resolve().as_uri())
        await page.wait_for_load_state("networkidle")
        await page.pdf(path=str(out), format="A4", print_background=True)
        await browser.close()


def main(args: Backup) -> int:
    frames = args.out / "frames"
    frames.mkdir(parents=True, exist_ok=True)
    for old in frames.glob("*.png"):
        old.unlink()
    status = capture(args, frames)
    slides = asyncio.run(slide_notes(args.url))
    grouped = frames_by_slide(frames)
    problems = []
    if status:
        problems.append(f"shoot.py exited with {status}: see frames/errors.txt and capture.log")
    missing = [s["id"] for s in slides if s["id"] not in grouped]
    problems += [f"no page for #{m}" for m in missing]
    pages = [f.path for s in slides if s["id"] in grouped for f in grouped[s["id"]]]
    blank = [p.name for p in pages if max(ImageStat.Stat(Image.open(p).convert("L")).stddev) < 2]
    problems += [f"blank frame {b}" for b in blank]
    slides_pdf(pages, args.out / "slides.pdf", args.jpeg_quality)
    slides_pptx([s for s in slides if s["id"] in grouped], grouped, args.out / "slides.pptx", args.jpeg_quality)
    asyncio.run(notes_pdf([s for s in slides if s["id"] in grouped], grouped, args.out / "notes.pdf"))
    layout = (frames / "layout.txt").read_text() if (frames / "layout.txt").exists() else ""
    report = [f"{len(slides)} slides, {len(pages)} pages in slides.pdf, {len(grouped)} slides in notes.pdf",
              f"slides.pdf {(args.out / 'slides.pdf').stat().st_size / 1e6:.1f} MB, "
              f"notes.pdf {(args.out / 'notes.pdf').stat().st_size / 1e6:.1f} MB, "
              f"slides.pptx {(args.out / 'slides.pptx').stat().st_size / 1e6:.1f} MB",
              *problems, "layout notes from the capture:", layout or "(none)"]
    (args.out / "check.txt").write_text("\n".join(report) + "\n")
    print("\n".join(report))
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main(tyro.cli(Backup)))
